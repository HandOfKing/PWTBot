"""Excel export: one match, or every match in a date range."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE = Font(bold=True, size=14)


def _mmss(s):
    if s is None: return None
    s = float(s); return f"{int(s // 60):02d}:{s % 60:05.2f}"


def _table(ws, headers, rows, start_row=1, flag_col=None, widths=None, freeze=True):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
        if flag_col is not None and r[flag_col]:
            for j in range(1, len(headers) + 1):
                ws.cell(row=i, column=j).fill = FLAG_FILL
    last = start_row + max(len(rows), 1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{last}"
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    for j, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or max([len(str(h))] + [len(str(r[j - 1])) for r in rows if r[j - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(j)].width = w if (widths or {}).get(h) else min(max(w, 8), 40)


EVENT_HEADERS = ["Date & time", "Round", "True time (s)", "True time", "Feed time (s)", "Feed delay (s)",
                 "Time source", "Type", "Killer", "Weapon", "Victim", "Victim team", "Confidence", "Flag", "Reviewed",
                 "Evidence image"]


def _event_rows(events):
    return [[e["recorded_at"], e["round_no"], e["true_time_s"], _mmss(e["true_time_s"]), e["feed_time_s"],
             e["feed_delay_s"], e["time_source"], e["event_type"], e["killer"], e["weapon"], e["victim"],
             e["victim_team"], e["confidence"], e["flag"], "yes" if e["reviewed"] else "", e["evidence_path"]]
            for e in events]


PLAYER_HEADERS = ["Player", "Nickname", "Team", "Eliminations", "Knocks", "Deaths", "Times knocked",
                  "Revives (inferred, at least)", "Damage dealt", "Board eliminations", "Board knock outs", "Check"]


def _player_rows(players):
    """Kill-feed counts next to the scoreboard's. "Board" is the end-of-match board when it was
    read (whole match, every player), else the sum of the round boards (partial in 6v6)."""
    rows = []
    for p in players:
        keys = p.keys()
        sb, sk = p["scoreboard_eliminations"], p["scoreboard_knock_outs"] if "scoreboard_knock_outs" in keys else None
        diffs = []
        if sb is not None and int(sb) != p["eliminations"]: diffs.append(f"elims {p['eliminations']} vs board {int(sb)}")
        if sk is not None and int(sk) != p["knocks"]: diffs.append(f"knocks {p['knocks']} vs board {int(sk)}")
        check = "" if sb is None and sk is None else ("; ".join(diffs) or "OK")
        rows.append([p["ign"], p["nickname"], p["team"], p["eliminations"], p["knocks"], p["deaths"],
                     p["times_knocked"], p["revives_inferred"], p["damage_dealt"], sb, sk, check])
    return rows


STAT_LABELS = {"eliminations": "Eliminations", "assists": "Assists", "damage_dealt": "Damage dealt",
               "damage_taken": "Damage taken", "knock_outs": "Knock outs"}
BOARD_ORDER = ("eliminations", "assists", "damage_dealt", "damage_taken", "knock_outs",
               "col_1", "col_2", "col_3", "col_4")
WRAP = Alignment(wrap_text=True, vertical="top")


def _scoreboard_sheet(ws, d):
    """The end-of-match board, every player and column, next to what PWT counted from the feed."""
    m = d["match"]
    ws["A1"] = f"Match {m['recorded_at']}"; ws["A1"].font = TITLE
    info = [("Source file", m["source_file"]), ("Room code", m["room_code"]), ("Team size", m["team_size"]),
            ("Rounds", m["rounds_played"]), ("Status", m["status"]), ("Processed", m["processed_at"]),
            ("Pipeline", m["pipeline_version"])]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(horizontal="left")
    board, names = {}, []
    for st in d["stats"]:
        if st["round_no"] is None:
            board.setdefault(st["ign"], {})[st["stat_name"]] = st["value"]
            if st["stat_name"] not in names: names.append(st["stat_name"])
    cols = [c for c in BOARD_ORDER if c in names] + [c for c in names if c not in BOARD_ORDER]
    r0 = len(info) + 4
    if not board:
        ws.cell(row=r0 - 1, column=1, value="The end-of-match scoreboard was not read in this recording; the "
                "board columns are empty. PWT's own counts are below.").font = Font(italic=True)
    rows, by_ign = [], {p["ign"]: p for p in d["players"]}
    for ign in list(by_ign) + [i for i in board if i not in by_ign]:       # every board row, even one with no team
        p, b, diffs = by_ign.get(ign), board.get(ign, {}), []
        if p is None:
            check = "on the board, but not in this match's player list"
        elif not board:
            check = ""
        elif not b:
            check = "not on the board"
        else:
            if b.get("eliminations") is not None and int(b["eliminations"]) != p["eliminations"]:
                diffs.append(f"elims {p['eliminations']} vs board {int(b['eliminations'])}")
            if b.get("knock_outs") is not None and int(b["knock_outs"]) != p["knocks"]:
                diffs.append(f"knocks {p['knocks']} vs board {int(b['knock_outs'])}")
            unread = [STAT_LABELS.get(c, c) for c in ("eliminations", "knock_outs") if b.get(c) is None]
            check = "; ".join(diffs) or ("OK" if not unread else "")
            if unread and not diffs:
                check = f"{' and '.join(unread)} unreadable on the board"
            elif unread:
                check += f" ({' and '.join(unread)} unreadable)"
        rows.append([ign, p["team"] if p else None] + [b.get(c) for c in cols]
                    + [p["eliminations"] if p else None, p["knocks"] if p else None, check])
    _table(ws, ["Player", "Team"] + [STAT_LABELS.get(c, c) for c in cols] + ["PWT eliminations", "PWT knocks", "Check"],
           rows, start_row=r0, freeze=False)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 22)


def _highlights_sheet(ws, d):
    from . import awards
    aw = awards.from_detail(d)
    ws["A1"] = "Match highlights"; ws["A1"].font = TITLE
    start = _notes_block(ws, aw.notes)
    rows = [[r.player, r.team, r.role, r.detail] for r in aw.roles]
    have = {r.player for r in aw.roles}
    for p in d["players"]:
        if p["ign"] not in have:
            n = aw.knock_counts.get(p["ign"], 0)
            rows.append([p["ign"], p["team"], "-", f"No role: {awards._plural(n, 'knock')} counted, no pattern strong enough."])
    rows.sort(key=lambda r: (r[1] or "", r[0].lower(), r[2] == "-"))
    _table(ws, ["Player", "Team", "Role", "What they did"], rows, start_row=start, widths={"What they did": 100}, freeze=False)
    for i in range(start + 1, start + 1 + len(rows)):
        ws.cell(row=i, column=4).alignment = WRAP


def _notes_block(ws, notes, width=4):
    for i, note in enumerate(notes, 2):
        c = ws.cell(row=i, column=1, value=note)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=width)
        ws.row_dimensions[i].height = 15 * (1 + len(note) // 120)
    return len(notes) + 3


def _players_sheet(ws, d):
    """Per player: knocks made and taken, and the full rankings of who they knocked / who knocked them."""
    from . import playerstats
    st = playerstats.from_detail(d)
    ws["A1"] = "Players: who knocked whom (knocks only)"; ws["A1"].font = TITLE
    start = _notes_block(ws, st.notes, width=6)
    rows = [[l.player, l.team, l.knocks, l.times_knocked,
             playerstats.ranking_text(l.knocked), playerstats.ranking_text(l.knocked_by)] for l in st.players]
    _table(ws, ["Player", "Team", "Knocks", "Times knocked", "Knocked most (ranking)", "Knocked by most (ranking)"],
           rows, start_row=start,
           widths={"Knocked most (ranking)": 70, "Knocked by most (ranking)": 70}, freeze=False)
    for i in range(start + 1, start + 1 + len(rows)):
        for j in (5, 6):
            ws.cell(row=i, column=j).alignment = WRAP


def _head_to_head_sheet(ws, d):
    """One row per player and opponent who knocked each other at least once: filter the Player column."""
    from . import playerstats
    st = playerstats.from_detail(d)
    _table(ws, ["Player", "Opponent", "Opponent team", "Knocked them", "Knocked by them"],
           [list(r) for r in st.pairs])


def _safe_sheet(wb, title, index, build, d, log_lines):
    """Build a sheet; a bug in it must not cost the whole result, so the error goes on the sheet and in the log."""
    ws = wb.create_sheet(title, index)
    try:
        build(ws, d)
    except Exception as exc:
        msg = f"{title} could not be built: {type(exc).__name__}: {exc}"
        wb.remove(ws); ws = wb.create_sheet(title, index)
        ws["A1"] = msg
        if log_lines is not None: log_lines.append(msg)


def export_match(conn, match_id, path, log_lines=None):
    """One workbook for a match: Scoreboard, Highlights, Players, Head to head, Events and, when given, the run log."""
    from .db import match_detail
    d = match_detail(conn, match_id)
    wb = Workbook()
    _scoreboard_sheet(wb.active, d); wb.active.title = "Scoreboard"
    _safe_sheet(wb, "Highlights", 1, _highlights_sheet, d, log_lines)
    _safe_sheet(wb, "Players", 2, _players_sheet, d, log_lines)
    _safe_sheet(wb, "Head to head", 3, _head_to_head_sheet, d, log_lines)
    _table(wb.create_sheet("Events"), EVENT_HEADERS, _event_rows(d["events"]), flag_col=13)
    if log_lines is not None:
        ws = wb.create_sheet("Run log")
        _table(ws, ["Run log"], [[l] for l in log_lines], widths={"Run log": 160})
    wb.save(path)
    return path


def export_range(conn, path, date_from=None, date_to=None):
    """Every match between two dates: match list, player totals, per-match player lines, all events."""
    from .db import list_matches
    matches = list_matches(conn, date_from, date_to)
    ids = [m["id"] for m in matches] or [-1]
    marks = ",".join("?" * len(ids))
    wb = Workbook()
    ws = wb.active; ws.title = "Matches"
    _table(ws, ["Date & time", "Room code", "Mode", "Team size", "Rounds", "Winner", "Players", "Eliminations",
                "Rows to review", "Status"],
           [[m["recorded_at"], m["room_code"], m["mode"], m["team_size"], m["rounds_played"], m["winner_team"],
             m["players"], m["eliminations"], m["to_review"], m["status"]] for m in matches])

    ws = wb.create_sheet("Player totals")
    rows = conn.execute(f"""SELECT ign, nickname, COUNT(DISTINCT match_id) AS matches, SUM(eliminations) AS el,
                                   SUM(knocks) AS kn, SUM(deaths) AS de, SUM(times_knocked) AS tk,
                                   SUM(revives_inferred) AS rv, SUM(damage_dealt) AS dmg
                            FROM v_player_match WHERE match_id IN ({marks})
                            GROUP BY player_id ORDER BY el DESC""", ids).fetchall()
    _table(ws, ["Player", "Nickname", "Matches", "Eliminations", "Knocks", "Deaths", "Times knocked",
                "Revives (inferred, at least)", "Damage dealt", "Elims per death"],
           [[r["ign"], r["nickname"], r["matches"], r["el"], r["kn"], r["de"], r["tk"], r["rv"], r["dmg"],
             round(r["el"] / max(r["de"], 1), 2)] for r in rows])

    ws = wb.create_sheet("Player per match")
    rows = conn.execute(f"SELECT * FROM v_player_match WHERE match_id IN ({marks}) ORDER BY recorded_at DESC, team, eliminations DESC", ids).fetchall()
    _table(ws, ["Date & time"] + PLAYER_HEADERS, [[r["recorded_at"]] + pr for r, pr in zip(rows, _player_rows(rows))])

    ws = wb.create_sheet("All events")
    ev = conn.execute(f"SELECT * FROM v_events WHERE match_id IN ({marks}) ORDER BY recorded_at, true_time_s", ids).fetchall()
    _table(ws, EVENT_HEADERS, _event_rows(ev), flag_col=13)
    wb.save(path)
    return path
