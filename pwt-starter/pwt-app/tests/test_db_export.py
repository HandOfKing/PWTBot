"""Run: python -m pytest tests   (or: python tests/test_db_export.py)"""
import os, sys, tempfile
from pathlib import Path
import _util  # noqa: F401  (puts the app root on sys.path)
from openpyxl import load_workbook
from pwt import db, export_excel


def _fresh():
    d = tempfile.mkdtemp()
    return db.connect(Path(d) / "pwt.db"), Path(d)


def _sample(conn):
    """One round, three feed rows, and that round's scoreboard -- the shape of a processed match."""
    players = [dict(ign="TheWolverine", team="blue"), dict(ign="PARABloodthirs", team="blue"),
               dict(ign="RGODxEMPEROR", team="red"), dict(ign="Makjets69", team="red")]
    rounds = [dict(round_no=1, start_s=2.0, end_s=13.46, winner_team="blue", blue_score=1, red_score=0)]
    e = dict(killer="TheWolverine", confidence=0.9, round_no=1)
    events = [dict(e, true_time_s=12.58, feed_time_s=12.58, time_source="feed (approx)", event_type="knock",
                   victim="RGODxEMPEROR", weapon="UMP45"),
              dict(e, true_time_s=13.25, feed_time_s=14.67, time_source="Remaining drop + helmet",
                   event_type="eliminated_knocked", victim="RGODxEMPEROR"),
              dict(e, true_time_s=13.25, feed_time_s=16.75, time_source="Remaining drop + helmet",
                   event_type="kill", victim="Makjets69", weapon="UMP45")]
    board = {"TheWolverine": (2, 200), "PARABloodthirs": (0, 0), "RGODxEMPEROR": (0, 52), "Makjets69": (0, 0)}
    stats = []
    for ign, (el, dmg) in board.items():
        stats += [dict(ign=ign, round_no=1, stat_name="eliminations", value=el),
                  dict(ign=ign, round_no=1, stat_name="damage_dealt", value=dmg)]
    match = dict(recorded_at="2026-10-01 22:53:00", source_file="sample.mkv", duration_s=21.5, mode="rounds",
                 team_size=6, rounds_played=1, winner_team="blue")
    return db.save_match(conn, match, players, rounds, events, stats)


def test_import_is_idempotent_and_reconciles():
    conn, _ = _fresh()
    a = _sample(conn)
    b = _sample(conn)
    assert a == b
    players = {p["ign"]: p for p in db.match_detail(conn, a)["players"]}
    w = players["TheWolverine"]
    assert (w["eliminations"], w["knocks"], w["deaths"]) == (2, 1, 0)
    assert players["RGODxEMPEROR"]["times_knocked"] == 1 and players["RGODxEMPEROR"]["deaths"] == 1
    for p in players.values():            # feed eliminations must match the round scoreboard
        assert p["eliminations"] == int(p["scoreboard_eliminations"])


def test_events_carry_true_and_feed_time():
    conn, _ = _fresh()
    mid = _sample(conn)
    ev = db.match_detail(conn, mid)["events"]
    kill = [e for e in ev if e["victim"] == "Makjets69"][0]
    assert kill["true_time_s"] == 13.25 and kill["feed_delay_s"] == 3.5 and kill["round_no"] == 1


def test_alias_and_list_filters():
    conn, _ = _fresh()
    mid = _sample(conn)
    assert db.get_or_create_player(conn, "TheWølverine") == db.get_or_create_player(conn, "TheWolverine")
    assert len(db.list_matches(conn, "2026-10-01", "2026-10-01")) == 1
    assert len(db.list_matches(conn, "2026-10-02")) == 0
    assert len(db.list_matches(conn, player="wolver")) == 1


def test_excel_exports():
    conn, d = _fresh()
    mid = _sample(conn)
    wb = load_workbook(export_excel.export_match(conn, mid, d / "m.xlsx"))
    assert wb.sheetnames == ["Summary", "Events", "Rounds", "Round scoreboards", "Match scoreboard"]
    assert wb["Events"].max_row == 4
    wb = load_workbook(export_excel.export_range(conn, d / "r.xlsx", "2026-10-01", "2026-10-31"))
    assert wb.sheetnames == ["Matches", "Player totals", "Player per match", "All events"]


def test_knock_outcomes_and_inferred_revives():
    """Synthetic 2-round match covering the revive rules in schema.sql (v_knock_outcomes)."""
    conn, _ = _fresh()
    P = lambda ign, team: dict(ign=ign, team=team)
    players = [P("A", "blue"), P("B", "blue"), P("Z", "blue"), P("X", "red"), P("Y", "red"), P("W", "red")]
    rounds = [dict(round_no=1, start_s=0, end_s=60), dict(round_no=2, start_s=70, end_s=130)]
    def ev(t, typ, k, v, rnd): return dict(true_time_s=t, feed_time_s=t, time_source="test", event_type=typ,
                                          killer=k, victim=v, round_no=rnd)
    events = [
        ev(10, "knock", "A", "X", 1),              # X knocked ...
        ev(14, "knock", "B", "X", 1),              # ... knocked again -> first knock = revived
        ev(16, "kill", "B", "X", 1),               # second knock -> eliminated
        ev(20, "knock", "A", "Y", 1),              # nothing more that round -> unresolved
        ev(80, "knock", "A", "X", 2),              # round 2: X knocked ...
        ev(85, "knock", "X", "B", 2),              # ... then X knocks B -> revived
        ev(90, "knock", "Z", "W", 2),              # Z knocks W
        ev(92, "knock", "Y", "Z", 2),              # Z gets knocked
        ev(93, "eliminated_knocked", "Z", "W", 2), # W dies on a team wipe credited to knocked Z: NOT proof Z is up
    ]
    mid = db.save_match(conn, dict(recorded_at="2026-10-02 21:00:00", source_file="synthetic.mkv"),
                        players, rounds, events)
    out = {(r["feed_time_s"]): r["outcome"] for r in conn.execute(
        "SELECT feed_time_s, outcome FROM v_knock_outcomes WHERE match_id=? ORDER BY feed_time_s", (mid,))}
    assert out == {10: "revived", 14: "eliminated", 20: "unresolved", 80: "revived",
                   85: "unresolved", 90: "eliminated", 92: "unresolved"}, out
    pm = {p["ign"]: p for p in db.match_detail(conn, mid)["players"]}
    assert (pm["X"]["times_knocked"], pm["X"]["revives_inferred"], pm["X"]["deaths"]) == (3, 2, 1)
    assert pm["Z"]["revives_inferred"] == 0 and pm["Z"]["eliminations"] == 1


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"): fn(); print("ok", name)
