"""pwt.playerstats -- who knocked whom, and revives by Chirag's rule, on a synthetic match (no recording needed)."""
import tempfile
from pathlib import Path
import _util  # noqa: F401  (puts the app root on sys.path)
from openpyxl import load_workbook
from pwt import db, export_excel, playerstats

BLUE = ["B1", "B2", "B3"]
RED = ["R1", "R2", "R3"]
TEAMS = {**{p: "blue" for p in BLUE}, **{p: "red" for p in RED}}


def ev(t, typ, killer, victim, rnd, flag=None, team="auto"):
    return dict(true_time_s=t, feed_time_s=t, time_source="test", event_type=typ, killer=killer, victim=victim,
                round_no=rnd, flag=flag, confidence=1.0,
                victim_team=TEAMS.get(victim) if team == "auto" else team)


def stats_of(events):
    """Save the events as a match and read the stats back through the real view (v_knock_outcomes)."""
    conn = db.connect(Path(tempfile.mkdtemp()) / "pwt.db")
    players = [dict(ign=p, team=t) for p, t in TEAMS.items()]
    rounds = [dict(round_no=r, start_s=r * 100, end_s=r * 100 + 90) for r in (1, 2, 3)]
    mid = db.save_match(conn, dict(recorded_at="2026-10-13 20:00:00", source_file="s.mkv", team_size=3),
                        players, rounds, events)
    return conn, mid, playerstats.from_detail(db.match_detail(conn, mid))


def line(st, p):
    return next(l for l in st.players if l.player == p)


def test_rankings_are_knocks_only_and_ordered():
    events = ([ev(110 + i, "knock", "B1", "R1", 1) for i in range(3)] + [ev(120, "knock", "B1", "R2", 1)]
              + [ev(130, "knock", "R1", "B1", 1)] + [ev(140, "kill", "B1", "R3", 1)]            # a kill is not a knock
              + [ev(150, "knock", "B1", "B2", 1)])                                              # a team-mate: ignored
    _, _, st = stats_of(events)
    b1, r1 = line(st, "B1"), line(st, "R1")
    assert b1.knocked == [("R1", 3), ("R2", 1)] and b1.knocks == 4
    assert b1.knocked_by == [("R1", 1)] and r1.knocked_by == [("B1", 3)]
    assert (("B1", "R1", "red", 3, 1) in st.pairs) and (("R1", "B1", "blue", 1, 3) in st.pairs)
    assert not any(p == "B1" and o == "B2" for p, o, *_ in st.pairs)


def test_revive_by_chirags_rule():
    events = [
        # round 1: R1 is knocked and never finished -> revived (survived)
        ev(110, "knock", "B1", "R1", 1),
        # round 2: R2 knocked twice then finished -> the first knock is a revive, the second was finished
        ev(210, "knock", "B1", "R2", 2), ev(220, "knock", "B3", "R2", 2), ev(225, "kill", "B3", "R2", 2),
        # round 3: B2 knocked, then knocks someone himself -> revived; later finished
        ev(310, "knock", "R3", "B2", 3), ev(320, "knock", "B2", "R1", 3), ev(330, "eliminated_knocked", "R3", "B2", 3),
    ]
    _, _, st = stats_of(events)
    r1, r2, b2 = line(st, "R1"), line(st, "R2"), line(st, "B2")
    assert (r1.times_knocked, r1.revived, r1.finished, r1.revived_unclear) == (2, 2, 0, 0)    # rounds 1 and 3, both survived
    assert r1.revived_confirmed == 0                                   # nobody saw him get up: only survived the round
    assert (r2.times_knocked, r2.revived, r2.finished, r2.revived_confirmed) == (2, 1, 1, 1)   # knocked again = confirmed
    assert (b2.times_knocked, b2.revived, b2.finished) == (1, 1, 0)        # acted again before dying: revived


def test_a_death_with_no_name_makes_a_survivor_unclear_not_revived():
    events = [ev(110, "knock", "B1", "R1", 1),
              ev(125, "kill", None, None, 1, flag="NO_FEED_ROW", team="red")]          # someone on red died, name unread
    _, _, st = stats_of(events)
    r1 = line(st, "R1")
    assert (r1.revived, r1.revived_unclear) == (0, 1)
    other_team = [ev(110, "knock", "B1", "R1", 1), ev(125, "kill", None, None, 1, flag="NO_FEED_ROW", team="blue")]
    assert (line(stats_of(other_team)[2], "R1").revived, line(stats_of(other_team)[2], "R1").revived_unclear) == (1, 0)


def test_times_knocked_counts_unnamed_knockers_and_the_workbook_has_both_sheets():
    events = [ev(110, "knock", None, "R1", 1), ev(120, "knock", "B1", "R1", 1)]          # zone / unread knocker + a named one
    conn, mid, st = stats_of(events)
    r1 = line(st, "R1")
    assert r1.times_knocked == 2 and r1.knocked_by == [("B1", 1)]
    d = Path(tempfile.mkdtemp())
    wb = load_workbook(export_excel.export_match(conn, mid, d / "p.xlsx"))
    rows = {r[0]: r for r in (tuple(c.value for c in row) for row in wb["Players"].iter_rows()) if r[0]}
    assert rows["R1"][3] == 2 and "B1 1" in rows["R1"][5]
    h2h = [tuple(c.value for c in row) for row in wb["Head to head"].iter_rows(min_row=2)]
    assert ("B1", "R1", "red", 1, 0) in h2h and ("R1", "B1", "blue", 0, 1) in h2h
